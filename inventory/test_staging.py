import csv
import io
import json
from datetime import date
from django.test import TestCase, Client
from django.contrib.auth import get_user_model
from django.core.files.uploadedfile import SimpleUploadedFile
from openpyxl import Workbook
from .models import Project, TabRecord, StagedInventoryRecord
from .specs import get_spec


class StagingTests(TestCase):
    def setUp(self):
        User = get_user_model()
        self.user = User.objects.create_user('stager', password='StrongPass!234')
        self.other = User.objects.create_user('other-stager', password='StrongPass!234')
        self.project = Project.objects.create(name='Staging project', code='staging')
        self.project.members.add(self.user, self.other)
        self.client.force_login(self.user)

    def upload(self, key, rows):
        stream = io.StringIO()
        writer = csv.DictWriter(stream, fieldnames=list(rows[0]))
        writer.writeheader()
        writer.writerows(rows)
        return self.client.post(f'/api/{key}/staging', {'file': SimpleUploadedFile('inventory.csv', stream.getvalue().encode(), 'text/csv')})

    def listing(self, key='sign'):
        return self.client.get(f'/api/{key}/staging').json()['records']

    def test_generic_sheet_names_accept_one_two_or_four_inventory_types(self):
        types = [('sign', 'SIGN_UID'), ('pavement', 'PM_UID'), ('lane', 'LANE_UID'), ('curb', 'CURB_UID')]
        for size in (1, 2, 4):
            with self.subTest(size=size):
                workbook = Workbook()
                workbook.active.title = 'Empty sheet'
                for number, (key, field) in enumerate(types[:size], 1):
                    sheet = workbook.create_sheet('Sheet' + str(number))
                    sheet.append([field])
                    sheet.append([str(size) + '-' + key])
                stream = io.BytesIO()
                workbook.save(stream)
                response = self.client.post('/api/curb/staging', {'file':SimpleUploadedFile('inventory.xlsx', stream.getvalue())})
                self.assertEqual(response.status_code, 201, response.content)
                self.assertEqual(response.json()['by_type'], {key:1 for key, field in types[:size]})
                self.assertEqual(response.json()['imported'], size)

    def test_repeat_uploads_and_duplicate_detection_for_every_type(self):
        for key, field in [('sign','SIGN_UID'), ('pavement','PM_UID'), ('lane','LANE_UID'), ('curb','CURB_UID')]:
            row = {field:'same', 'STREET_NAME':'Original'}
            self.assertEqual(self.upload(key, [row]).json()['imported'], 1)
            repeated = self.upload(key, [row]).json()
            self.assertEqual(repeated['imported'], 0)
            self.assertEqual(repeated['skipped'], 1)
            self.assertEqual(len(self.listing(key)), 1)
            self.upload(key, [{**row, 'STREET_NAME':'Changed'}])
            rows = self.listing(key)
            self.assertTrue(all(r['duplicate'] for r in rows))
            self.assertEqual(self.approve(rows, key).status_code, 400)
            self.client.delete(f'/api/{key}/staging/{rows[1]["id"]}', json.dumps({'version':rows[1]['version']}), content_type='application/json')
            self.assertEqual(self.approve(self.listing(key), key).status_code, 200)
            self.upload(key, [row])
            self.assertTrue(self.listing(key)[0]['live_duplicate'])

    def test_approval_and_rejection_are_separate_for_each_inventory_type(self):
        for key, field in [('sign','SIGN_UID'), ('pavement','PM_UID'), ('lane','LANE_UID'), ('curb','CURB_UID')]:
            self.assertEqual(self.upload(key, [{field:key}]).status_code, 201)
        pavement = self.listing('pavement')
        self.assertEqual(self.approve(pavement, 'sign').status_code, 409)
        self.assertEqual(self.approve(self.listing('sign'), 'sign').status_code, 200)
        self.assertEqual(list(TabRecord.objects.values_list('tab', flat=True)), ['sign'])
        self.assertEqual(set(StagedInventoryRecord.objects.values_list('tab', flat=True)), {'pavement','lane','curb'})
        lane = self.listing('lane')[0]
        response = self.client.post('/api/lane/staging/reject', json.dumps({'rows':[{'id':lane['id'], 'version':lane['version']}]}), content_type='application/json')
        self.assertEqual(response.status_code, 200)
        self.assertEqual(set(StagedInventoryRecord.objects.values_list('tab', flat=True)), {'pavement','curb'})
        self.assertEqual(self.approve(pavement, 'pavement').status_code, 200)
        self.assertEqual(set(TabRecord.objects.values_list('tab', flat=True)), {'sign','pavement'})
        self.assertEqual(StagedInventoryRecord.objects.get().tab, 'curb')

    def test_multi_inventory_workbook_routes_all_sheets_atomically(self):
        workbook = Workbook()
        workbook.remove(workbook.active)
        for title, header in [('SIGN_&_RETRO', 'SIGN_UID'), ('PAVEMENT', 'PM_UID'), ('LANE', 'LANE_UID'), ('CURB', 'CURB_UID')]:
            sheet = workbook.create_sheet(title)
            sheet.append([header])
            sheet.append([title + '-1'])
        workbook['CURB'].cell(1, 2, 'CURB_MARKING_TYPE')
        workbook['CURB'].cell(2, 2, 'Painted')
        workbook['PAVEMENT'].cell(2, 2, 'CAR COVERED')
        def upload():
            stream = io.BytesIO()
            workbook.save(stream)
            return self.client.post('/api/lane/staging', {'file': SimpleUploadedFile('combined.xlsx', stream.getvalue())})
        response = upload()
        self.assertEqual(response.status_code, 201, response.content)
        self.assertEqual(response.json()['by_type'], {'sign':1, 'pavement':1, 'lane':1, 'curb':1})
        self.assertEqual(TabRecord.objects.count(), 0)
        self.assertEqual(StagedInventoryRecord.objects.get(tab='curb').data['CURB_MARKING'], 'Painted')
        self.assertEqual(StagedInventoryRecord.objects.get(tab='pavement').data['COMMENT'], 'CAR COVERED')
        self.assertTrue(all(row.owner == self.user and row.project == self.project for row in StagedInventoryRecord.objects.all()))
        workbook['CURB'].cell(1, 3, 'TYPO_COLUMN')
        workbook['CURB'].cell(2, 3, 'invalid')
        response = upload()
        self.assertEqual(response.status_code, 201, response.content)
        self.assertEqual(response.json()['imported'], 0)
        self.assertEqual(StagedInventoryRecord.objects.count(), 4)

    def test_extra_columns_are_ignored_for_all_inventory_types(self):
        for key, field in [('sign', 'SIGN_UID'), ('pavement', 'PM_UID'), ('lane', 'LANE_UID'), ('curb', 'CURB_UID')]:
            response = self.upload(key, [{field: 'extra-test', 'SMALL_IMAGE_LINK': 'small.jpg', 'NEW_COLUMN': 'extra'}])
            self.assertEqual(response.status_code, 201, response.content)
            row = StagedInventoryRecord.objects.get(tab=key)
            self.assertEqual(row.data[field], 'extra-test')
            self.assertNotIn('SMALL_IMAGE_LINK', row.data)
            self.assertNotIn('NEW_COLUMN', row.data)

    def test_duplicate_extra_headers_and_extra_only_rows_are_ignored(self):
        content = b'SIGN_UID,SMALL_IMAGE_LINK,SMALL_IMAGE_LINK\nA,one,two\n,extra,only\n'
        response = self.client.post('/api/sign/staging', {'file': SimpleUploadedFile('inventory.csv', content, 'text/csv')})
        self.assertEqual(response.status_code, 201, response.content)
        self.assertEqual(response.json()['imported'], 1)
        self.assertEqual(StagedInventoryRecord.objects.get().data['SIGN_UID'], 'A')

    def test_duplicate_known_headers_still_fail_without_importing(self):
        content = b'SIGN_UID,SIGN UID,SMALL_IMAGE_LINK\nA,B,extra\n'
        response = self.client.post('/api/sign/staging', {'file': SimpleUploadedFile('inventory.csv', content, 'text/csv')})
        self.assertEqual(response.status_code, 400)
        self.assertEqual(StagedInventoryRecord.objects.count(), 0)

    def approve(self, rows, key='sign'):
        return self.client.post(f'/api/{key}/staging/approve', json.dumps({'rows':[{'id':r['id'], 'version':r['version']} for r in rows]}), content_type='application/json')

    def test_bulk_reject_is_scoped_atomic_and_preserves_live_records(self):
        self.upload('sign', [{'SIGN_UID':'A'}, {'SIGN_UID':'B'}, {'SIGN_UID':'C'}])
        rows = self.listing()
        def reject(selected, client=None):
            return (client or self.client).post('/api/sign/staging/reject', json.dumps({'rows': selected}), content_type='application/json')
        selected = [{'id':r['id'], 'version':r['version']} for r in rows[:2]]
        self.assertEqual(reject([selected[0], dict(selected[1], version='stale')]).status_code, 409)
        self.assertEqual(StagedInventoryRecord.objects.count(), 3)
        other_client = Client()
        other_client.force_login(self.other)
        self.assertEqual(reject(selected, other_client).status_code, 409)
        self.assertEqual(StagedInventoryRecord.objects.count(), 3)
        self.assertEqual(self.approve([rows[2]]).status_code, 200)
        self.assertEqual(reject(selected).json(), {'rejected': 2})
        self.assertEqual(StagedInventoryRecord.objects.count(), 0)
        self.assertEqual(TabRecord.objects.count(), 1)
        self.assertEqual(reject([]).status_code, 400)
        csrf_client = Client(enforce_csrf_checks=True)
        csrf_client.force_login(self.user)
        self.assertEqual(reject(selected, csrf_client).status_code, 403)

    def test_normalization_and_staging_only(self):
        response = self.upload('sign', [{'SIGN_UID':'A', 'SIGN_CONDITION':'?', 'NIGHT_SIMULATION':'NV', 'SHARED_POLE':'Yes ', 'SHEETING_MATERIAL':'NA', 'LEVEL(R)':'NA', 'INSP_DATE':'2026-09-14'}])
        self.assertEqual(response.status_code, 201, response.content)
        self.assertEqual(TabRecord.objects.count(), 0)
        stage = StagedInventoryRecord.objects.get()
        self.assertEqual(stage.inspection_date, date(2026,9,14))
        self.assertEqual(stage.data['SHARED_POLE'],'Yes')
        self.assertEqual(stage.data['SHEETING_MATERIAL'],'None')
        self.assertEqual(stage.data['LEVEL(R)'],'')
        self.assertEqual(stage.data['_IMPORTED_NIGHT_SIMULATION'],'N')
        self.assertEqual(stage.data['Night Sumulation Visibility'],'No')
        self.assertEqual(self.listing()[0]['status'],'Flagged: Condition unclear (?)')
        self.assertEqual(self.approve(self.listing()).status_code,400)

    def test_duplicates_require_explicit_rejection(self):
        self.assertEqual(self.upload('sign',[{'SIGN_UID':'same','SIGN_CONDITION':'Good'},{'SIGN_UID':'same','SIGN_CONDITION':'Fair'}]).status_code,201)
        rows=self.listing()
        self.assertTrue(all(r['duplicate'] for r in rows))
        self.assertEqual(self.approve(rows).status_code,400)
        self.assertEqual(TabRecord.objects.count(),0)
        self.assertEqual(self.client.delete(f"/api/sign/staging/{rows[1]['id']}", json.dumps({'version':rows[1]['version']}), content_type='application/json').status_code,200)
        self.assertEqual(StagedInventoryRecord.objects.count(),1)
        self.assertEqual(self.approve(self.listing()).status_code,200)
        self.assertEqual(StagedInventoryRecord.objects.count(),0)
        self.assertEqual(TabRecord.objects.get().data['SIGN_CONDITION'],'Good')
        self.upload('sign',[{'SIGN_UID':'same'}])
        self.assertTrue(self.listing()[0]['live_duplicate'])
        self.assertEqual(self.approve(self.listing()).status_code,400)

    def test_edit_and_optimistic_lock(self):
        self.upload('sign',[{'SIGN_UID':'unique','SIGN_CONDITION':'?'}])
        row=self.listing()[0]
        data=dict(row['data']);data['SIGN_CONDITION']='Good'
        response=self.client.put(f"/api/sign/staging/{row['id']}",json.dumps({'row':data,'version':row['version']}),content_type='application/json')
        self.assertEqual(response.status_code,200)
        self.assertEqual(self.approve([row]).status_code,409)
        self.assertEqual(self.approve(self.listing()).status_code,200)

    def test_all_four_types_and_bulk(self):
        for key in ('sign','pavement','lane','curb'):
            columns=get_spec(key)['columns']
            field=next(c for c in columns if c not in ('ID','INSP_DATE'))
            self.assertEqual(self.upload(key,[{field:'one'},{field:'two'}]).status_code,201)
            self.assertEqual(self.approve(self.listing(key),key).status_code,200)
            self.assertEqual(TabRecord.objects.filter(tab=key).count(),2)

    def test_access_control_csrf_and_bad_file(self):
        self.upload('sign',[{'SIGN_UID':'private'}])
        row=self.listing()[0]
        self.client.force_login(self.other)
        self.assertEqual(self.listing(),[])
        self.assertEqual(self.approve([row]).status_code,409)
        self.assertEqual(self.client.delete(f"/api/sign/staging/{row['id']}",json.dumps({'version':row['version']}),content_type='application/json').status_code,404)
        csrf=Client(enforce_csrf_checks=True);csrf.force_login(self.user)
        self.assertEqual(csrf.post('/api/sign/staging/approve','{}',content_type='application/json').status_code,403)
        self.assertEqual(self.client.post('/api/sign/staging',{'file':SimpleUploadedFile('bad.xlsx',b'bad')}).status_code,400)
        self.assertEqual(self.upload('sign',[{'INSP_DATE':'bad'}]).status_code,400)
        self.assertEqual(StagedInventoryRecord.objects.count(),1)

    def test_xlsx_native_dates(self):
        wb=Workbook();wb.active.append(['SIGN_UID','INSP_DATE']);wb.active.append(['xlsx',date(2026,9,14)])
        stream=io.BytesIO();wb.save(stream)
        response=self.client.post('/api/sign/staging',{'file':SimpleUploadedFile('sheet.xlsx',stream.getvalue())})
        self.assertEqual(response.status_code,201,response.content)
        self.assertEqual(StagedInventoryRecord.objects.get().inspection_date,date(2026,9,14))

    def test_section_id_alias_is_import_only(self):
        from openpyxl import load_workbook
        for key in ('sign', 'pavement', 'lane', 'curb'):
            response = self.upload(key, [{'SECTION ID': '1008-260'}])
            self.assertEqual(response.status_code, 201, response.content)
            staged = self.listing(key)
            self.assertEqual(staged[0]['data']['ST_ID'], '1008-260')
            self.assertNotIn('SECTION ID', staged[0]['data'])
            self.assertEqual(self.approve(staged, key).status_code, 200)
            response = self.client.get(f'/api/{key}/export')
            self.assertEqual(response.status_code, 200)
            workbook = load_workbook(io.BytesIO(response.content))
            headers = [cell.value for cell in workbook.active[1]]
            self.assertNotIn('SECTION ID', headers)
            self.assertEqual(workbook.active.cell(2, headers.index('ST_ID') + 1).value, '1008-260')
            workbook.close()

    def test_xlsx_section_id_header(self):
        workbook = Workbook()
        workbook.active.append(['SECTION ID', 'SIGN_UID'])
        workbook.active.append(['1008-260', 'SR_1008-260_P1_S1'])
        stream = io.BytesIO()
        workbook.save(stream)
        response = self.client.post('/api/sign/staging', {'file': SimpleUploadedFile('inventory.xlsx', stream.getvalue())})
        self.assertEqual(response.status_code, 201, response.content)
        self.assertEqual(self.listing()[0]['data']['ST_ID'], '1008-260')

    def test_night_image_link_and_legacy_visibility_are_separate(self):
        response = self.upload('sign', [{'SECTION ID':'S', 'SIGN_UID':'link-test', 'NIGHT_SIMULATION':'NV', 'NIGHT_SIMULATION_LINK':'https://stpaul.bluedometech.com/night/test.png'}])
        self.assertEqual(response.status_code, 201, response.content)
        row = self.listing()[0]
        self.assertEqual(row['data']['NIGHT_SIMULATION'], 'https://stpaul.bluedometech.com/night/test.png')
        self.assertEqual(row['data']['Night Sumulation Visibility'], 'No')
        self.assertEqual(row['data']['_IMPORTED_NIGHT_SIMULATION'], 'N')
        self.assertEqual(self.approve([row]).status_code, 200)
        self.assertEqual(TabRecord.objects.get().data['NIGHT_SIMULATION'], row['data']['NIGHT_SIMULATION'])

    def test_night_image_excel_column_round_trip(self):
        from .views import write_sheet
        link = 'https://stpaul.bluedometech.com/night/example.png'
        workbook = Workbook()
        write_sheet('sign', workbook.active, [{'SIGN_UID':'night-excel', 'NIGHT_SIMULATION':link, 'Night Sumulation Visibility':'Yes'}])
        headers = [cell.value for cell in workbook.active[1]]
        self.assertEqual(headers.count('Night Simulation Image'), 1)
        self.assertIn('Night Sumulation Visibility', headers)
        self.assertEqual(workbook.active.cell(2, headers.index('Night Simulation Image') + 1).hyperlink.target, link)
        stream = io.BytesIO()
        workbook.save(stream)
        response = self.client.post('/api/sign/staging', {'file': SimpleUploadedFile('night-images.xlsx', stream.getvalue())})
        self.assertEqual(response.status_code, 201, response.content)
        row = self.listing()[0]
        self.assertEqual(row['data']['NIGHT_SIMULATION'], link)
        self.assertEqual(row['data']['Night Sumulation Visibility'], 'Yes')
        self.assertEqual(self.approve([row]).status_code, 200)
        self.assertEqual(TabRecord.objects.get().data['NIGHT_SIMULATION'], link)

    def test_zero_values_survive_import(self):
        response = self.upload('sign', [{'SIGN_UID':'zero','LATITUDE':0,'LONGITUDE':0,'LEVEL(R)':0}])
        self.assertEqual(response.status_code, 201)
        data = self.listing()[0]['data']
        for field in ('LATITUDE', 'LONGITUDE', 'LEVEL(R)'):
            self.assertEqual(data[field], '0')

    def test_unsafe_image_urls_and_unlabelled_data_are_rejected(self):
        for link in ('javascript:alert(1)', 'file:///c:/secret.png', '//evil.example/image.png', 'https://user:password@example.com/image.png'):
            response = self.upload('sign', [{'SIGN_UID':'unsafe','NIGHT_SIMULATION_LINK':link}])
            self.assertEqual(response.status_code, 400, response.content)
        response = self.client.post('/api/sign/staging', {'file':SimpleUploadedFile('bad.csv', b'SIGN_UID,\nuid,unlabelled\n')})
        self.assertEqual(response.status_code, 400)
        self.assertEqual(StagedInventoryRecord.objects.count(), 0)

    def test_configuration_labels_cannot_break_out_of_script(self):
        from .models import InventorySection
        section = InventorySection.objects.get(pk='sign')
        config = dict(section.configuration)
        config['tab_label'] = '</script><script>alert(1)</script>'
        section.configuration = config
        section.save()
        response = self.client.get('/')
        self.assertEqual(response.status_code, 200)
        self.assertNotContains(response, '</script><script>alert(1)</script>')
        self.assertContains(response, 'inventory-tabs-data')

    def test_trusted_image_origins_are_exposed_consistently(self):
        from django.test import override_settings
        with override_settings(INVENTORY_IMAGE_ORIGINS=['https://images.example.com']):
            response = self.client.get('/api/sign/spec')
            self.assertEqual(response.json()['image_origins'], ['https://images.example.com'])
            self.assertIn('https://images.example.com', response['Content-Security-Policy'])
            self.assertNotIn('img-src *', response['Content-Security-Policy'])

    def test_optional_supplied_workbook_import(self):
        import os
        filename = os.environ.get('INVENTORY_REVIEW_WORKBOOK')
        if not filename:
            self.skipTest('Set INVENTORY_REVIEW_WORKBOOK for the real-sheet integration check.')
        from openpyxl import load_workbook
        workbook = load_workbook(filename, read_only=True, data_only=True)
        expected_rows = sum(any(value not in (None, '') for value in row) for sheet in workbook for row in list(sheet.values)[1:])
        workbook.close()
        with open(filename, 'rb') as stream:
            response = self.client.post('/api/sign/staging', {'file':stream})
        self.assertEqual(response.status_code, 201, response.content)
        self.assertEqual(response.json()['imported'], expected_rows)
        self.assertEqual(StagedInventoryRecord.objects.count(), expected_rows)
        self.assertEqual(TabRecord.objects.count(), 0)
        rows = self.listing()
        self.assertTrue(all(row['data']['ST_ID'] for row in rows))
        clean = [row for row in rows if row['status'] == 'OK'][:3]
        self.assertEqual(self.approve(clean).status_code, 200)
        self.assertEqual(TabRecord.objects.count(), len(clean))
        self.assertTrue(all(not record.data['NIGHT_SIMULATION'] for record in TabRecord.objects.all()))

    def test_import_does_not_query_section_for_every_row(self):
        from django.db import connection
        from django.test.utils import CaptureQueriesContext
        with CaptureQueriesContext(connection) as queries:
            response = self.upload('sign', [{'SIGN_UID':f'performance-{i}'} for i in range(100)])
        self.assertEqual(response.status_code, 201)
        config_queries = [query for query in queries if 'inventory_inventorysection' in query['sql']]
        self.assertLess(len(config_queries), 15)
