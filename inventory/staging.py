"""Review-only imports. Live records are created exclusively by approve()."""
import csv
import io
import json
from itertools import chain
from collections import Counter
from contextlib import ExitStack
from datetime import date, datetime
from urllib.parse import urlsplit

from django.conf import settings
from django.core.exceptions import ValidationError
from django.db import transaction
from django.db.models import Max
from django.http import JsonResponse
from django.views.decorators.http import require_http_methods, require_POST
from openpyxl import load_workbook

from .concurrency import section_write_locks as locks
from .forms import validate_xlsx_upload
from .models import StagedInventoryRecord, TabRecord
from .specs import TAB_ORDER, get_spec
from .views import api_login_required, require_selected_project, parse_json_body, next_tab_id

VISIBILITY = "Night Sumulation Visibility"


def scope(request, key):
    if key not in TAB_ORDER:
        return None, None, JsonResponse({"error": "Unknown inventory type."}, status=404)
    project, error = require_selected_project(request)
    if error:
        return None, None, error
    rows = StagedInventoryRecord.objects.filter(project=project, tab=key)
    if not request.user.is_staff:
        rows = rows.filter(owner=request.user)
    return project, rows, None


def normalize(key, values, columns=None):
    if columns is None:
        columns = get_spec(key)["columns"]
    data = {column: str(values.get(column) if values.get(column) is not None else "").strip() for column in columns}
    if values.get("_IMPORTED_NIGHT_SIMULATION") in {"Y", "N"}:
        data["_IMPORTED_NIGHT_SIMULATION"] = values["_IMPORTED_NIGHT_SIMULATION"]
    data["SHARED_POLE"] = data.get("SHARED_POLE", "").strip()
    if data.get("SHEETING_MATERIAL", "").upper() in {"NONE", "NA"}:
        data["SHEETING_MATERIAL"] = "None"
    if data.get("LEVEL(R)", "").upper() == "NA":
        data["LEVEL(R)"] = ""
    night = data.get("NIGHT_SIMULATION", "")
    if night.upper() in {"N", "NV", "NA", "Y"}:
        data["_IMPORTED_NIGHT_SIMULATION"] = "Y" if night.upper() == "Y" else "N"
        data[VISIBILITY] = "Yes" if night.upper() == "Y" else "No"
        data["NIGHT_SIMULATION"] = ""
    link = str(values.get("_NIGHT_SIMULATION_LINK", "") or "").strip()
    if link:
        if data.get("NIGHT_SIMULATION") and data["NIGHT_SIMULATION"] != link:
            raise ValueError("Conflicting night simulation image links.")
        data["NIGHT_SIMULATION"] = link
    for field in ("IMAGE_LINK", "NIGHT_SIMULATION"):
        value = data.get(field, "")
        if not value:
            continue
        try:
            parsed_url = urlsplit(value)
            valid = (parsed_url.scheme in {"http", "https"} and bool(parsed_url.hostname) and not parsed_url.username and not parsed_url.password) or (not parsed_url.scheme and not parsed_url.netloc and value.startswith("/uploads/images/"))
        except ValueError:
            valid = False
        if not valid:
            raise ValueError(f"{field} must be an HTTP(S) image link or an application upload path.")
    raw_date = values.get("INSP_DATE", "")
    parsed = None
    if isinstance(raw_date, datetime):
        parsed = raw_date.date()
    elif isinstance(raw_date, date):
        parsed = raw_date
    elif str(raw_date or "").strip():
        for fmt in ("%m-%d-%Y", "%Y-%m-%d", "%m/%d/%Y", "%m/%d/%y", "%Y/%m/%d", "%d-%b-%Y", "%d %b %Y"):
            try:
                parsed = datetime.strptime(str(raw_date).strip(), fmt).date()
                break
            except ValueError:
                pass
        if parsed is None:
            raise ValueError("INSP_DATE must be a valid date (MM-DD-YYYY, MM/DD/YYYY or YYYY-MM-DD).")
    if "INSP_DATE" in columns:
        data["INSP_DATE"] = parsed.strftime("%m-%d-%Y") if parsed else ""
    if data.get(VISIBILITY, "") not in {"", "Yes", "No"}:
        raise ValueError("Night Simulation Visibility must be Yes, No or blank.")
    return data, parsed


UID_FIELDS = {"sign": "SIGN_UID", "pavement": "PM_UID", "lane": "LANE_UID", "curb": "CURB_UID"}


def fingerprint(data):
    return json.dumps({k: v for k, v in data.items() if k != "ID" and v not in (None, "")}, sort_keys=True, default=str)


def uid(data, key="sign"):
    value = str(data.get(UID_FIELDS[key], "") or "").strip().casefold()
    return value or "row:" + fingerprint(data)


def flags_for(project, key):
    staged = Counter(uid(data, key) for data in StagedInventoryRecord.objects.filter(project=project, tab=key).values_list("data", flat=True))
    live = {uid(data, key) for data in TabRecord.objects.filter(project=project, tab=key).values_list("data", flat=True)}
    return staged, live


def row_payload(row, counts, live):
    duplicate = bool(uid(row.data, row.tab) and (counts[uid(row.data, row.tab)] > 1 or uid(row.data, row.tab) in live))
    unclear = row.tab == "sign" and row.data.get("SIGN_CONDITION", "").strip() == "?"
    flags = []
    if duplicate:
        flags.append("Flagged: Duplicate — needs decision")
    if unclear:
        flags.append("Flagged: Condition unclear (?)")
    return {"id": row.pk, "data": row.data, "status": "; ".join(flags) or "OK", "duplicate": duplicate,
            "unclear": unclear, "live_duplicate": uid(row.data, row.tab) in live, "source": row.source_file,
            "source_row": row.source_row, "version": str(row.revision), "duplicate_key": uid(row.data, row.tab), "uid_field": UID_FIELDS[row.tab]}


def parse_sheet(iterator, key, project, owner, source):
    columns = get_spec(key)["columns"]
    headers = [str(value or "").strip() for value in next(iterator, [])]
    aliases = {column.upper().replace("_", " "): column for column in get_spec(key)["columns"]}
    aliases["NIGHT SIMULATION"] = "NIGHT_SIMULATION"
    if key == "sign":
        for name in ("NIGHT SIMULATION LINK", "NIGHT SIMULATION IMAGE", "NIGHT SIMULATION IMAGE LINK"):
            aliases[name] = "_NIGHT_SIMULATION_LINK"
        aliases["NIGHT SIMULATION VISIBILITY"] = VISIBILITY
    if "ST_ID" in get_spec(key)["columns"]:
        aliases["SECTION ID"] = "ST_ID"
    if key == "curb":
        aliases["CURB MARKING TYPE"] = "CURB_MARKING"
    headers = [aliases.get(h.upper().replace("_", " "), h) for h in headers]
    # Legacy pavement/lane sheets sometimes leave their final notes header blank.
    if "COMMENT" in columns and "COMMENT" not in headers and headers and headers[-1] == "" and headers.count("") == 1:
        headers[-1] = "COMMENT"
    known = set(get_spec(key)["columns"])
    if key == "sign":
        known.add("_NIGHT_SIMULATION_LINK")
    recognized_headers = [h for h in headers if h in known]
    if not headers or len(set(recognized_headers)) != len(recognized_headers):
        raise ValueError("The sheet has missing or duplicate column headers.")
    if not any(h in known - {"ID"} for h in headers):
        raise ValueError("No inventory columns were recognized. Use the selected inventory's Excel export headers.")
    # Extra source columns are ignored; only configured inventory fields are saved.
    pending = []
    for line, values in enumerate(iterator, 2):
        if line > 50001:
            raise ValueError("The sheet exceeds the 50,000 scanned-row limit, including blank rows.")
        if not any(value not in (None, "") for value in values):
            continue
        if any(value not in (None, "") for value in values[len(headers):]) or any(not header and index < len(values) and values[index] not in (None, "") for index, header in enumerate(headers)):
            raise ValueError(f"Row {line}: data appears in a column without a header.")
        if len(pending) >= 10000:
            raise ValueError("Upload at most 10,000 rows per sheet.")
        raw = {h: values[index] if index < len(values) else "" for index, h in enumerate(headers) if h in known}
        if not any(value not in (None, "") for value in raw.values()):
            continue
        try:
            data, inspection_date = normalize(key, raw, columns=columns)
        except ValueError as exc:
            raise ValueError(f"Row {line}: {exc}") from exc
        pending.append(StagedInventoryRecord(project=project, owner=owner, tab=key, data=data,
                       inspection_date=inspection_date, source_file=source[:255], source_row=line))
    if not pending:
        raise ValueError("The sheet contains no inventory rows.")
    return pending


@api_login_required
@require_http_methods(["GET", "POST"])
def staging(request, key):
    project, rows, error = scope(request, key)
    if error:
        return error
    if request.method == "GET":
        counts, live = flags_for(project, key)
        return JsonResponse({"records": [row_payload(row, counts, live) for row in rows], "columns": get_spec(key)["columns"]})
    upload = request.FILES.get("file")
    if upload is None:
        return JsonResponse({"error": "Choose a CSV or XLSX inventory sheet."}, status=400)
    if upload.size > settings.MAX_XLSX_UPLOAD_BYTES:
        return JsonResponse({"error": "Inventory sheet exceeds the upload size limit."}, status=400)
    workbook = None
    try:
        if upload.name.lower().endswith(".xlsx"):
            validate_xlsx_upload(upload)
            workbook = load_workbook(upload, read_only=True, data_only=True)
            names = {get_spec(section)["export_sheet_name"].casefold(): section for section in TAB_ORDER}
            names.update({section: section for section in TAB_ORDER})
            names["sign_&_retro"] = "sign"
            sheets = []
            for sheet in workbook:
                iterator = iter(sheet.iter_rows(values_only=True))
                headers = next(iterator, ())
                # Ignore genuinely empty worksheets, but never silently drop data.
                first_data = None
                for line, values in enumerate(iterator, 2):
                    if line > 50001:
                        raise ValueError(f"{sheet.title}: The sheet exceeds the 50,000 scanned-row limit.")
                    if any(value not in (None, "") for value in values):
                        first_data = (line, values)
                        break
                if first_data is None:
                    continue
                normalized = {str(value or "").strip().upper().replace(" ", "_") for value in headers}
                indicators = {
                    "sign": {"SIGN_UID", "POLE_ID", "SIGN_CONDITION", "SHEETING_MATERIAL"},
                    "pavement": {"PM_UID", "PM_ID"},
                    "lane": {"LANE_UID", "LM_ID"},
                    "curb": {"CURB_UID", "CM_ID", "CURB_MARKING", "CURB_MARKING_TYPE"},
                }
                detected = [section for section, fields in indicators.items() if normalized & fields]
                sheet_key = names.get(sheet.title.strip().casefold())
                if len(detected) == 1:
                    sheet_key = detected[0]
                elif len(detected) > 1:
                    raise ValueError(f"{sheet.title}: Columns contain multiple inventory types. Keep each type on a separate worksheet.")
                if sheet_key is None:
                    sheet_key = key
                line, values = first_data
                sheets.append((sheet_key, sheet.title, chain([headers], [()] * (line - 2), [values], iterator)))
            if not sheets:
                raise ValueError("The workbook contains no inventory rows.")
        elif upload.name.lower().endswith(".csv"):
            sheets = [(key, upload.name, iter(csv.reader(io.StringIO(upload.read().decode("utf-8-sig")))))]
        else:
            raise ValueError("Upload an XLSX or UTF-8 CSV file.")
        pending = []
        counts = Counter()
        for sheet_key, title, iterator in sheets:
            try:
                parsed = parse_sheet(iterator, sheet_key, project, request.user, upload.name)
            except ValueError as exc:
                raise ValueError(f"{title}: {exc}") from exc
            pending.extend(parsed)
            counts[sheet_key] += len(parsed)
        with ExitStack() as stack:
            for section in TAB_ORDER:
                if section in counts:
                    stack.enter_context(locks[section])
            with transaction.atomic():
                existing = {
                    section: {fingerprint(data) for data in StagedInventoryRecord.objects.filter(
                        project=project, tab=section, owner=request.user
                    ).values_list("data", flat=True)}
                    for section in counts
                }
                skipped = Counter()
                fresh = []
                for row in pending:
                    if fingerprint(row.data) in existing[row.tab]:
                        skipped[row.tab] += 1
                    else:
                        fresh.append(row)
                pending = fresh
                counts = Counter(row.tab for row in pending)
                StagedInventoryRecord.objects.bulk_create(pending)
    except (ValueError, ValidationError, UnicodeError, csv.Error) as exc:
        return JsonResponse({"error": str(exc)}, status=400)
    except Exception:
        import logging
        logging.getLogger(__name__).exception("Inventory staging import failed")
        return JsonResponse({"error": "The workbook could not be read. Check that it is a valid inventory sheet."}, status=400)
    finally:
        if workbook is not None:
            workbook.close()
    return JsonResponse({"imported": len(pending), "by_type": dict(counts), "skipped": sum(skipped.values()), "skipped_by_type": dict(skipped)}, status=201)


@api_login_required
@require_http_methods(["PUT", "DELETE"])
def staging_detail(request, key, row_id):
    project, rows, error = scope(request, key)
    if error:
        return error
    with locks[key], transaction.atomic():
        row = rows.select_for_update().filter(pk=row_id).first()
        if row is None:
            return JsonResponse({"error": "Staged row not found."}, status=404)
        body = parse_json_body(request)
        if body.get("version") != str(row.revision):
            return JsonResponse({"error": "This staged row changed. Refresh Staging and try again."}, status=409)
        if request.method == "DELETE":
            row.delete()
            return JsonResponse({"ok": True})
        if not isinstance(body.get("row"), dict):
            return JsonResponse({"error": "Provide a row object."}, status=400)
        try:
            row.data, row.inspection_date = normalize(key, body["row"])
        except ValueError as exc:
            return JsonResponse({"error": str(exc)}, status=400)
        row.revision += 1
        row.save(update_fields=["data", "inspection_date", "updated_at", "revision"])
    return JsonResponse({"ok": True})


@api_login_required
@require_POST
def reject(request, key):
    project, rows, error = scope(request, key)
    if error:
        return error
    selected = parse_json_body(request).get("rows")
    if not isinstance(selected, list) or not selected or len(selected) > 10000 or any(not isinstance(item, dict) or type(item.get("id")) is not int for item in selected):
        return JsonResponse({"error": "Select between 1 and 10,000 staged rows to reject."}, status=400)
    versions = {item["id"]: item.get("version") for item in selected}
    with locks[key], transaction.atomic():
        records = list(rows.select_for_update().filter(pk__in=versions))
        if len(records) != len(versions) or any(versions[row.pk] != str(row.revision) for row in records):
            return JsonResponse({"error": "Some rows changed, were removed, or are not accessible. Refresh Staging and try again. No rows were rejected."}, status=409)
        rows.filter(pk__in=versions).delete()
    return JsonResponse({"rejected": len(records)})


@api_login_required
@require_POST
def approve(request, key):
    project, rows, error = scope(request, key)
    if error:
        return error
    selected = parse_json_body(request).get("rows")
    if not isinstance(selected, list) or not selected or len(selected) > 10000 or any(not isinstance(item, dict) or type(item.get("id")) is not int for item in selected):
        return JsonResponse({"error": "Select staged rows to approve."}, status=400)
    versions = {item["id"]: item.get("version") for item in selected}
    with locks[key], transaction.atomic():
        records = list(rows.select_for_update().filter(pk__in=versions))
        if len(records) != len(versions):
            return JsonResponse({"error": "Some rows no longer exist or are not accessible."}, status=409)
        counts, live = flags_for(project, key)
        for row in records:
            if versions[row.pk] != str(row.revision):
                return JsonResponse({"error": "Some staged rows changed. Refresh before approving."}, status=409)
            payload = row_payload(row, counts, live)
            if payload["status"] != "OK":
                return JsonResponse({"error": "Resolve flagged rows before approval. No selected rows were moved."}, status=400)
        next_id = next_tab_id(key)
        order = TabRecord.objects.filter(project=project, tab=key).aggregate(n=Max("display_order"))["n"] or 0
        for offset, row in enumerate(records, 1):
            data = dict(row.data)
            data.pop("ID", None)
            TabRecord.objects.create(project=project, owner=row.owner, tab=key, tab_record_id=next_id + offset - 1,
                                     display_order=order + offset, data=data)
        rows.filter(pk__in=versions).delete()
    return JsonResponse({"approved": len(records)})
