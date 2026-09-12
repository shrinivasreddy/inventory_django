from pathlib import Path

from django.core.files.base import ContentFile
from django.core.management.base import BaseCommand, CommandError
from django.db import transaction
from django.utils.text import slugify
from openpyxl import load_workbook

from inventory.models import Jurisdiction, MutcdReference


class Command(BaseCommand):
    help = "Import MUTCD code, description, and embedded images from an Excel workbook."

    def add_arguments(self, parser):
        parser.add_argument("workbook", type=Path)
        parser.add_argument("--jurisdiction", required=True, help="Jurisdiction code, for example us-ca")

    def handle(self, *args, **options):
        workbook_path = options["workbook"].resolve()
        if not workbook_path.is_file() or workbook_path.suffix.lower() != ".xlsx":
            raise CommandError("Choose an existing .xlsx workbook.")
        try:
            jurisdiction = Jurisdiction.objects.get(code=options["jurisdiction"])
        except Jurisdiction.DoesNotExist as exc:
            raise CommandError("The selected jurisdiction does not exist.") from exc

        workbook = load_workbook(workbook_path, data_only=True)
        created = updated = images = skipped = 0
        with transaction.atomic():
            for worksheet in workbook.worksheets:
                images_by_row = {
                    image.anchor._from.row + 1: image for image in worksheet._images
                }
                for row_number in range(3, worksheet.max_row + 1):
                    code = str(worksheet.cell(row_number, 1).value or "").strip()
                    description = str(worksheet.cell(row_number, 3).value or "").strip()
                    source_sheet = str(worksheet.cell(row_number, 5).value or "").strip()
                    if not code or not description:
                        skipped += 1
                        continue
                    reference, was_created = MutcdReference.objects.update_or_create(
                        jurisdiction=jurisdiction,
                        mutcd_code=code,
                        description=description,
                        defaults={"source_sheet": source_sheet},
                    )
                    created += int(was_created)
                    updated += int(not was_created)
                    embedded_image = images_by_row.get(row_number)
                    if embedded_image is not None:
                        extension = (embedded_image.format or "png").lower()
                        if extension == "jpeg":
                            extension = "jpg"
                        if reference.image:
                            reference.image.delete(save=False)
                        filename = f"{slugify(code) or 'mutcd'}-{reference.pk}.{extension}"
                        reference.image.save(
                            filename, ContentFile(embedded_image._data()), save=True
                        )
                        images += 1

        self.stdout.write(
            self.style.SUCCESS(
                f"Imported {created + updated} MUTCD references for {jurisdiction}: "
                f"{created} created, {updated} updated, {images} images, {skipped} skipped."
            )
        )
